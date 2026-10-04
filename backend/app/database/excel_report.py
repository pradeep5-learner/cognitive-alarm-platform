import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="6C56E8", end_color="6C56E8", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="19152E")


def style_header_row(ws, row_num, num_cols):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left")


def autosize_columns(ws):
    for col_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 3, 12), 45)


def generate_excel_report(data):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Overview"
    ws1["A1"] = "Cognitive Alarm Platform — Habit Report"
    ws1["A1"].font = TITLE_FONT
    ws1["A2"] = f"{data['user_name']} ({data['user_email']})"
    ws1["A3"] = f"Generated {data['generated_at']}"
    overview_rows = [
        ["Metric", "Value"],
        ["Total Alarms", data["total_alarms"]],
        ["Total Wake-Up Sessions", data["total_sessions"]],
        ["Verified Sessions", data["verified_sessions"]],
        ["Current Streak (days)", data["streak"]["current_streak"]],
        ["Longest Streak (days)", data["streak"]["longest_streak"]],
    ]
    for r in overview_rows:
        ws1.append(r)
    style_header_row(ws1, 5, 2)
    autosize_columns(ws1)

    ws2 = wb.create_sheet("Habit Score")
    hs = data["habit_scores"]
    habit_rows = [
        ["Component", "Score (%)"],
        ["Wake-Up Consistency", hs["wake_consistency_score"]],
        ["Challenge Completion", hs["challenge_completion_score"]],
        ["Snooze Reduction", hs["snooze_reduction_score"]],
        ["Sleep Schedule Adherence", hs["sleep_adherence_score"]],
        ["Productivity", hs["productivity_score"]],
        ["Total Habit Score", hs["total_score"]],
    ]
    for r in habit_rows:
        ws2.append(r)
    style_header_row(ws2, 1, 2)
    autosize_columns(ws2)

    ws3 = wb.create_sheet("Behavioral Analytics")
    ws3.append(["Day", "Avg Snoozes", "Avg Attempts", "Sessions"])
    style_header_row(ws3, 1, 4)
    for day_stat in data["behavior"]["day_of_week_breakdown"]:
        ws3.append([day_stat["day"], day_stat["avg_snoozes"], day_stat["avg_attempts"], day_stat["sessions"]])
    ws3.append([])
    ws3.append(["Best Day", data["behavior"]["best_day"]])
    ws3.append(["Worst Day", data["behavior"]["worst_day"]])
    ws3.append(["Avg Response Time (min)", data["behavior"]["avg_response_time_minutes"]])
    autosize_columns(ws3)

    ws4 = wb.create_sheet("Sleep & Productivity")
    ws4.append(["Sleep Duration (h)", data["sleep"]["sleep_duration_hours"]])
    ws4.append(["Sleep Category", data["sleep"]["category"]])
    ws4.append(["Sleep Insight", data["sleep"]["insight"]])
    ws4.append([])
    prod = data["productivity"]
    ws4.append(["Productivity Correlation", prod["correlation"]])
    ws4.append(["Correlation Strength", prod["strength"]])
    ws4.append(["Days Analyzed", prod["days_analyzed"]])
    ws4.append(["Productivity Insight", prod["insight"]])
    autosize_columns(ws4)

    ws5 = wb.create_sheet("Challenge Performance")
    ws5.append(["Challenge Type", "Avg Attempts", "Sessions Completed", "Skill Weight"])
    style_header_row(ws5, 1, 4)
    for ctype, stats in data["challenge_perf"].items():
        ws5.append([
            ctype.replace("_", " ").title(),
            stats["avg_attempts"] if stats["avg_attempts"] is not None else "Not tried",
            stats["count"],
            stats["skill_weight"]
        ])
    autosize_columns(ws5)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer