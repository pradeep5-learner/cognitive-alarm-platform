import io
from openpyxl import Workbook
from app.database.excel_report import style_header_row, autosize_columns, TITLE_FONT


def _header(ws, values):
    ws.append(values)
    style_header_row(ws, ws.max_row, len(values))


def generate_system_excel(data):
    wb = Workbook()
    stats = data["stats"]
    rec = data["recommendations"]

    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "Cognitive Alarm Platform — System Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {data['generated_at']}"
    ws.append([])
    _header(ws, ["Metric", "Value"])
    for row in [
        ["Total accounts", stats["total_users"]],
        ["Total alarms", stats["total_alarms"]],
        ["Active alarms", stats["active_alarms"]],
        ["Wake-up sessions", stats["total_wakeup_sessions"]],
        ["Platform completion rate (%)", stats["platform_completion_rate"]],
        ["Average habit score", stats["avg_habit_score"]],
        ["Notifications sent", data["system"]["notifications"]],
        ["Habit score snapshots", data["system"]["habit_score_snapshots"]],
        ["Productivity ratings logged", data["system"]["productivity_ratings"]],
    ]:
        ws.append(row)
    ws.append([])
    _header(ws, ["Role", "Accounts"])
    for r in data["roles"]:
        ws.append([r["role"], r["count"]])
    ws.append([])
    _header(ws, ["Difficulty level", "Users"])
    for level, count in stats["difficulty_distribution"].items():
        ws.append([level.title(), count])
    autosize_columns(ws)

    ws2 = wb.create_sheet("Users")
    _header(ws2, ["ID", "Name", "Email", "Role", "Coach", "Sessions", "Completion (%)", "Habit Score"])
    for u in data["users"]:
        ws2.append([u["id"], u["name"], u["email"], u["role"], u["coach"], u["sessions"], u["completion_rate"], u["habit_score"]])
    autosize_columns(ws2)

    ws3 = wb.create_sheet("Activity")
    _header(ws3, ["Date", "Sessions", "Completed", "Active users"])
    for a in data["activity"]:
        ws3.append([a["date"], a["sessions"], a["completed"], a["active_users"]])
    autosize_columns(ws3)

    ws4 = wb.create_sheet("Challenges")
    _header(ws4, ["Challenge type", "Sessions", "Completed", "Avg attempts"])
    for c in data["challenge_usage"]:
        ws4.append([c["challenge_type"].replace("_", " ").title(), c["sessions"], c["completed"], c["avg_attempts"]])
    autosize_columns(ws4)

    ws5 = wb.create_sheet("Recommendations")
    _header(ws5, ["Priority", "Count"])
    for p in ("high", "medium", "low"):
        ws5.append([p.title(), rec["by_priority"][p]])
    ws5.append([])
    _header(ws5, ["Category", "Count"])
    for c in rec["by_category"]:
        ws5.append([c["category"].replace("_", " ").title(), c["count"]])
    ws5.append([])
    _header(ws5, ["Most common recommendation", "Users"])
    for r in rec["top_recommendations"]:
        ws5.append([r["title"], r["count"]])
    ws5.append([])
    _header(ws5, ["User", "Email", "Top priority", "Recommendations", "Titles"])
    for u in rec["users"]:
        ws5.append([u["name"], u["email"], u["top_priority"].title(), u["count"], "; ".join(u["titles"])])
    autosize_columns(ws5)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer